"""
Utilidades de entrada/salida para PIAPIME.

Funciones para leer/escribir CSV, Excel y validar datos dosis-respuesta.
"""

from __future__ import annotations

import os
from typing import Optional

import numpy as np
import pandas as pd


# ---------------------------------------------------------------------------
# Lectura de datos
# ---------------------------------------------------------------------------

def read_csv(path: str) -> pd.DataFrame:
    """Lee un archivo CSV intentando detectar el separador automáticamente.

    Retorna un DataFrame con al menos dos columnas numéricas.
    Lanza ValueError si no puede parsear el archivo.
    """
    errors = []
    for sep in [",", ";", "\t", " "]:
        try:
            df = pd.read_csv(path, sep=sep, decimal=".", engine="python",
                             skip_blank_lines=True)
            if df.shape[1] >= 2:
                # Intentar convertir columnas a numéricas
                df_num = df.apply(pd.to_numeric, errors="coerce")
                if df_num.iloc[:, :2].notna().sum().min() >= 2:
                    return df_num.dropna(how="all")
        except Exception as e:
            errors.append(str(e))

    raise ValueError(f"No se pudo leer el CSV '{os.path.basename(path)}'. "
                     f"Verifica el formato del archivo.")


def read_excel(path: str, sheet: Optional[str] = None) -> pd.DataFrame:
    """Lee un archivo Excel.

    Si `sheet` es None, lee la primera hoja.
    Retorna un DataFrame con columnas numéricas.
    """
    try:
        if sheet is not None:
            df = pd.read_excel(path, sheet_name=sheet, engine="openpyxl")
        else:
            df = pd.read_excel(path, sheet_name=0, engine="openpyxl")

        df_num = df.apply(pd.to_numeric, errors="coerce")
        return df_num.dropna(how="all")
    except Exception as e:
        raise ValueError(f"No se pudo leer el Excel '{os.path.basename(path)}': {e}")


def get_excel_sheets(path: str) -> list[str]:
    """Retorna la lista de hojas en un archivo Excel."""
    try:
        xl = pd.ExcelFile(path, engine="openpyxl")
        return xl.sheet_names
    except Exception:
        return []


# ---------------------------------------------------------------------------
# Validación de datos
# ---------------------------------------------------------------------------

def validate_data(x: np.ndarray, y: np.ndarray) -> tuple[bool, str]:
    """Valida un par de arreglos dosis-respuesta.

    Retorna (True, "") si los datos son válidos,
    o (False, mensaje_error) si hay algún problema.
    """
    if len(x) == 0 or len(y) == 0:
        return False, "Los arreglos de datos están vacíos."
    if len(x) != len(y):
        return False, f"Las longitudes no coinciden: x={len(x)}, y={len(y)}."
    if not np.any(np.isfinite(x) & (x > 0)):
        return False, "No hay valores de concentración válidos (deben ser > 0)."
    if not np.any(np.isfinite(y)):
        return False, "No hay valores de respuesta válidos."
    valid = np.isfinite(x) & np.isfinite(y) & (x > 0)
    n_valid = int(np.sum(valid))
    if n_valid < 4:
        return False, (f"Se necesitan al menos 4 puntos válidos para el ajuste "
                       f"(encontrados: {n_valid}).")
    return True, ""


# ---------------------------------------------------------------------------
# Exportación de resultados
# ---------------------------------------------------------------------------

def _result_to_rows(dataset_name: str, result) -> list[dict]:
    """Convierte un FittingResult a una lista de diccionarios para exportar."""
    rows = []

    # Fila de resumen
    rows.append({
        "Dataset": dataset_name,
        "Modelo": result.model_name,
        "R²": round(result.r_squared, 6) if np.isfinite(result.r_squared) else "",
        "RMSE": round(result.rmse, 6) if np.isfinite(result.rmse) else "",
        "AIC": round(result.aic, 4) if np.isfinite(result.aic) else "",
        "BIC": round(result.bic, 4) if np.isfinite(result.bic) else "",
        "EC50": result.ec50 if np.isfinite(result.ec50) else "",
        "EC50_CI_Inferior": result.ec50_ci[0] if np.isfinite(result.ec50_ci[0]) else "",
        "EC50_CI_Superior": result.ec50_ci[1] if np.isfinite(result.ec50_ci[1]) else "",
        "Tipo": "Resumen",
        "Parámetro": "",
        "Valor": "",
        "Error_Estándar": "",
        "CI_95_Inferior": "",
        "CI_95_Superior": "",
    })

    # Filas de parámetros
    for pname, pval in result.parameters.items():
        perr = result.param_errors.get(pname, float("nan"))
        ci = result.ci_95.get(pname, (float("nan"), float("nan")))
        rows.append({
            "Dataset": dataset_name,
            "Modelo": result.model_name,
            "R²": "",
            "RMSE": "",
            "AIC": "",
            "BIC": "",
            "EC50": "",
            "EC50_CI_Inferior": "",
            "EC50_CI_Superior": "",
            "Tipo": "Parámetro",
            "Parámetro": pname,
            "Valor": round(pval, 8) if np.isfinite(pval) else "",
            "Error_Estándar": round(perr, 8) if np.isfinite(perr) else "",
            "CI_95_Inferior": round(ci[0], 8) if np.isfinite(ci[0]) else "",
            "CI_95_Superior": round(ci[1], 8) if np.isfinite(ci[1]) else "",
        })

    return rows


def export_results_csv(results: dict, path: str) -> None:
    """Exporta todos los resultados de ajuste a un archivo CSV.

    Args:
        results: {dataset_name: FittingResult}
        path: ruta del archivo de salida.
    """
    all_rows = []
    for dname, result in results.items():
        all_rows.extend(_result_to_rows(dname, result))

    if not all_rows:
        raise ValueError("No hay resultados para exportar.")

    df = pd.DataFrame(all_rows)
    df.to_csv(path, index=False, encoding="utf-8-sig")


def export_results_excel(results: dict, path: str) -> None:
    """Exporta todos los resultados de ajuste a un archivo Excel con formato.

    Args:
        results: {dataset_name: FittingResult}
        path: ruta del archivo de salida (.xlsx).
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Eliminar hoja por defecto

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    summary_fill = PatternFill("solid", fgColor="2E4B8A")
    summary_font = Font(color="FFFFFF", bold=True)
    alt_fill = PatternFill("solid", fgColor="EBF0FA")
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for dname, result in results.items():
        # Nombre de hoja (máx 31 caracteres, sin caracteres especiales)
        sheet_name = dname[:31].replace("/", "-").replace("\\", "-").replace("*", "")
        ws = wb.create_sheet(title=sheet_name)

        # --- Resumen ---
        ws.append(["RESUMEN DEL AJUSTE"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.merge_cells("A1:D1")

        headers_summary = ["Parámetro de Ajuste", "Valor"]
        ws.append(headers_summary)
        for cell in ws[2]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        summary_data = [
            ("Dataset", dname),
            ("Modelo", result.model_name),
            ("R²", f"{result.r_squared:.6f}" if np.isfinite(result.r_squared) else "N/A"),
            ("RMSE", f"{result.rmse:.6f}" if np.isfinite(result.rmse) else "N/A"),
            ("AIC", f"{result.aic:.4f}" if np.isfinite(result.aic) else "N/A"),
            ("BIC", f"{result.bic:.4f}" if np.isfinite(result.bic) else "N/A"),
            ("EC50", f"{result.ec50:.6e}" if np.isfinite(result.ec50) else "N/A"),
            ("EC50 IC 95% Inferior", f"{result.ec50_ci[0]:.6e}" if np.isfinite(result.ec50_ci[0]) else "N/A"),
            ("EC50 IC 95% Superior", f"{result.ec50_ci[1]:.6e}" if np.isfinite(result.ec50_ci[1]) else "N/A"),
        ]
        for i, (lbl, val) in enumerate(summary_data):
            ws.append([lbl, val])
            row_idx = i + 3
            for cell in ws[row_idx]:
                cell.border = border
            if i % 2 == 0:
                for cell in ws[row_idx]:
                    cell.fill = alt_fill

        ws.append([])

        # --- Parámetros ---
        param_row = ws.max_row + 1
        ws.append(["PARÁMETROS DEL MODELO"])
        ws.cell(row=param_row, column=1).font = Font(bold=True, size=12)
        ws.merge_cells(f"A{param_row}:E{param_row}")

        headers_params = ["Parámetro", "Valor", "Error Estándar", "IC 95% Inferior", "IC 95% Superior"]
        ws.append(headers_params)
        hrow = ws.max_row
        for cell in ws[hrow]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = border

        for i, (pname, pval) in enumerate(result.parameters.items()):
            perr = result.param_errors.get(pname, float("nan"))
            ci = result.ci_95.get(pname, (float("nan"), float("nan")))
            row_data = [
                pname,
                round(pval, 8) if np.isfinite(pval) else "N/A",
                round(perr, 8) if np.isfinite(perr) else "N/A",
                round(ci[0], 8) if np.isfinite(ci[0]) else "N/A",
                round(ci[1], 8) if np.isfinite(ci[1]) else "N/A",
            ]
            ws.append(row_data)
            rrow = ws.max_row
            for cell in ws[rrow]:
                cell.border = border
            if i % 2 == 0:
                for cell in ws[rrow]:
                    cell.fill = alt_fill

        # Ajustar ancho de columnas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    wb.save(path)
